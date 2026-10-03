import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from docx_helpers import (
    set_cell_background,
    add_footer_page_number,
    add_header_footer_to_section
)

# Colors
COLOR_NAVY = RGBColor(31, 73, 125)       # #1F497D Primary
COLOR_DARK = RGBColor(30, 41, 59)        # #1E293B Body text
COLOR_MUTED = RGBColor(100, 116, 139)    # #64748B Secondary
COLOR_ACCENT = RGBColor(14, 116, 144)    # #0E7490 Teal/Cyan accent
COLOR_RED = RGBColor(185, 28, 28)        # #B91C1C Red accent

HEX_HEADER_BG = "EBF2FA"                 # Soft elegant ice blue
HEX_ROW_ALT = "F8FAFC"
HEX_CODE_BG = "F1F5F9"

def setup_document():
    doc = Document()
    
    # Configure A4 & HUIT Margins: Top 3.5cm, Bottom 3.0cm, Left 3.5cm, Right 2.0cm
    section = doc.sections[0]
    section.page_width = Inches(8.27)     # 210mm
    section.page_height = Inches(11.69)   # 297mm
    section.top_margin = Inches(1.38)     # 35mm
    section.bottom_margin = Inches(1.18)  # 30mm
    section.left_margin = Inches(1.38)    # 35mm
    section.right_margin = Inches(0.79)   # 20mm
    
    # Set default style to Times New Roman
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = COLOR_DARK
    style_normal.paragraph_format.line_spacing = 1.3
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.space_before = Pt(0)
    
    return doc

def add_cover_page(doc):
    # Front cover section
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r1 = p.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    r2 = p.add_run("TRƯỜNG ĐẠI HỌC CÔNG THƯƠNG TP. HỒ CHÍ MINH (HUIT)\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(13)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_NAVY
    
    r3 = p.add_run("KHOA CÔNG NGHỆ THÔNG TIN – BỘ MÔN AN TOÀN THÔNG TIN\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True
    
    # Decorative line
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(36)
    r_line = p_line.add_run("━━━━━━━ ★ ━━━━━━━")
    r_line.font.name = "Times New Roman"
    r_line.font.size = Pt(11)
    r_line.font.color.rgb = COLOR_NAVY

    # Project Type
    p_type = doc.add_paragraph()
    p_type.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_type.paragraph_format.space_after = Pt(18)
    r_type = p_type.add_run("BÁO CÁO TIẾN ĐỘ ĐỒ ÁN CHUYÊN NGÀNH\nNGÀNH AN TOÀN THÔNG TIN")
    r_type.font.name = "Times New Roman"
    r_type.font.size = Pt(14)
    r_type.font.bold = True
    r_type.font.color.rgb = COLOR_NAVY

    # Title Box
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.space_before = Pt(18)
    r_title_label = p_title.add_run("ĐỀ TÀI:\n")
    r_title_label.font.name = "Times New Roman"
    r_title_label.font.size = Pt(12)
    r_title_label.font.italic = True
    
    r_title = p_title.add_run("XÂY DỰNG MÔ HÌNH KIỂM THỬ LỖ HỔNG SMB TRÊN WINDOWS BẰNG KALI LINUX")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(40)
    r_sub = p_sub.add_run("HỒ SƠ HỌC THUẬT TOÀN VĂN: CHƯƠNG 1 VÀ CHƯƠNG 2")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_ACCENT

    # Info Table (Borderless layout table)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    labels = [
        "Giảng viên hướng dẫn:",
        "Sinh viên thực hiện 1:",
        "Sinh viên thực hiện 2:",
        "Sinh viên thực hiện 3:",
        "Chuyên ngành đào tạo:",
        "Khóa học & Niên khóa:"
    ]
    values = [
        "ThS. Ngô Quốc Huy",
        "Lâm Gia Bảo  –  MSSV: 2033216350  –  Lớp: 12DHBM03",
        "Nguyễn Minh Thắng  –  MSSV: 2033216558  –  Lớp: 12DHBM03",
        "Nguyễn Hoài Tiến  –  MSSV: 2033216575  –  Lớp: 12DHBM09",
        "An toàn thông tin (Mã ngành: 7480202)",
        "Đại học chính quy Khóa 12 (2023 – 2027)"
    ]
    
    col_widths = [Inches(2.4), Inches(3.6)]
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(labels[row_idx])
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                run.font.bold = True
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(values[row_idx])
                run.font.name = "Times New Roman"
                run.font.size = Pt(12)
                if row_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = COLOR_NAVY

    # Location & Year
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(60)
    p_foot.paragraph_format.space_after = Pt(0)
    r_foot = p_foot.add_run("TP. HỒ CHÍ MINH, NĂM 2026")
    r_foot.font.name = "Times New Roman"
    r_foot.font.size = Pt(12)
    r_foot.font.bold = True

    doc.add_page_break()

def add_preliminary_section(doc):
    # TOC & Lists
    h_toc = doc.add_heading("MỤC LỤC TỔNG QUÁT", level=1)
    h_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h_toc.paragraph_format.space_before = Pt(12)
    h_toc.paragraph_format.space_after = Pt(18)
    for r in h_toc.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

    toc_items = [
        ("CHƯƠNG 1. CƠ SỞ LÝ THUYẾT VÀ CÔNG CỤ KIỂM THỬ SMB", "Trang 1"),
        ("  1.1. Tổng quan về giao thức SMB", "Trang 1"),
        ("    1.1.1. Khái niệm và vai trò của SMB trong hệ điều hành Windows", "Trang 1"),
        ("    1.1.2. Mô hình Client – Server", "Trang 2"),
        ("    1.1.3. Phân biệt SMBv1, SMBv2 và SMBv3", "Trang 3"),
        ("    1.1.4. Phân tích cổng mạng TCP 139 và TCP 445", "Trang 4"),
        ("    1.1.5. Quy trình trao đổi bản tin cơ bản", "Trang 5"),
        ("  1.2. Phân tích nhóm lỗ hổng MS17-010 và mã độc WannaCry", "Trang 6"),
        ("    1.2.1. Bối cảnh lịch sử và tác động an ninh", "Trang 6"),
        ("    1.2.2. Danh mục CVE trong bản tin MS17-010 và phạm vi của đề tài", "Trang 7"),
        ("    1.2.3. Phân tích cơ chế phát sinh lỗi FEA trong CVE-2017-0144", "Trang 8"),
        ("    1.2.4. Bản vá chính thức KB4012212 và cơ chế khắc phục", "Trang 9"),
        ("  1.3. Công cụ kiểm thử và kịch bản dò quét", "Trang 10"),
        ("    1.3.1. Kali Linux trong môi trường lab kiểm thử", "Trang 10"),
        ("    1.3.2. Nmap và kịch bản smb-vuln-ms17-010.nse", "Trang 11"),
        ("    1.3.3. Metasploit Framework và module ms17_010_eternalblue", "Trang 13"),
        ("  1.4. Khung đánh giá 4 cấp độ và giải pháp phòng thủ tổng thể", "Trang 14"),
        ("    1.4.1. Khung đánh giá trạng thái an ninh SMB theo 4 cấp độ", "Trang 14"),
        ("    1.4.2. Các giải pháp phòng thủ chiều sâu cho giao thức SMB", "Trang 15"),
        ("  TỔNG KẾT CHƯƠNG 1", "Trang 17"),
        ("  TÀI LIỆU THAM KHẢO CHƯƠNG 1", "Trang 18"),
        ("CHƯƠNG 2. THIẾT KẾ VÀ TRIỂN KHAI MÔ HÌNH THỰC NGHIỆM", "Trang 20"),
        ("  2.1. Yêu cầu và nguyên tắc thiết kế mô hình thực nghiệm", "Trang 20"),
        ("    2.1.1. Nguyên tắc đạo đức nghề nghiệp và phạm vi kiểm thử được ủy quyền", "Trang 20"),
        ("    2.1.2. Yêu cầu về tính cô lập và kiểm soát rủi ro mạng", "Trang 20"),
        ("    2.1.3. Phương pháp luận nghiên cứu chuỗi nhân quả (Causal-Chain)", "Trang 21"),
        ("  2.2. Thiết kế kiến trúc mạng lab doanh nghiệp mô phỏng", "Trang 22"),
        ("    2.2.1. Cấu trúc mạng phân đoạn 3 VLAN", "Trang 22"),
        ("    2.2.2. Nguyên lý Inter-VLAN Routing và Firewall Inspection", "Trang 23"),
        ("    2.2.3. Sơ đồ Topo mạng kiến trúc doanh nghiệp mô phỏng", "Trang 23"),
        ("    2.2.4. Phân bổ không gian địa chỉ IP và thông số kỹ thuật", "Trang 24"),
        ("  2.3. Mô hình máy trạng thái đơn biến (Single-Variable Causal State Machine)", "Trang 25"),
        ("    2.3.1. Thiết kế chuỗi 4 trạng thái Snapshot trên cùng một hệ thống mục tiêu", "Trang 25"),
        ("    2.3.2. Vai trò đối chứng nghiệp vụ dương tính của trạm Client tại VLAN 30", "Trang 27"),
        ("    2.3.3. Phương pháp luận chứng minh tính bất khả thi của đường tấn công", "Trang 28"),
        ("  2.4. Phương pháp luận và chuẩn hóa kỹ thuật theo mô hình bạch hộp", "Trang 30"),
        ("    2.4.1. Kỹ thuật 1: Khảo sát tiếp cận và nhận diện dịch vụ (TCP SYN Scan)", "Trang 30"),
        ("    2.4.2. Kỹ thuật 2: Thăm dò khả năng đàm phán phương ngữ (Dialect Probe)", "Trang 31"),
        ("    2.4.3. Kỹ thuật 3: Dò quét phát hiện dấu hiệu an toàn bằng Nmap NSE", "Trang 32"),
        ("    2.4.4. Kỹ thuật 4: Khảo sát đối chứng độc lập bằng Metasploit Auxiliary", "Trang 33"),
        ("    2.4.5. Kỹ thuật 5: Xác minh mức độ tác động có kiểm soát (Metasploit Exploit)", "Trang 34"),
        ("  2.5. Ma trận bằng chứng đa nguồn và quy trình vận hành an toàn", "Trang 36"),
        ("    2.5.1. Thiết kế Ma trận bằng chứng thu thập đa nguồn (Evidence Matrix)", "Trang 36"),
        ("    2.5.2. Quy trình kiểm chứng lưu lượng đa điểm (Multi-Point PCAP)", "Trang 37"),
        ("    2.5.3. Kế hoạch ứng phó sự cố màn hình xanh (BSOD) và Rollback", "Trang 38"),
        ("  TỔNG KẾT CHƯƠNG 2", "Trang 39"),
        ("  TÀI LIỆU THAM KHẢO CHƯƠNG 2", "Trang 40"),
    ]

    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.line_spacing = 1.25
    p_toc.paragraph_format.space_after = Pt(2)
    for title, pg in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(1)
        r_t = p.add_run(title)
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(11.5)
        if title.startswith("CHƯƠNG") or title.startswith("TỔNG KẾT") or title.startswith("TÀI LIỆU"):
            r_t.font.bold = True
            r_t.font.color.rgb = COLOR_NAVY
        elif title.startswith("  1.") or title.startswith("  2."):
            r_t.font.bold = True
        
        # Dots
        dots_count = max(5, 75 - len(title))
        r_dots = p.add_run(" " + "·" * dots_count + " ")
        r_dots.font.name = "Times New Roman"
        r_dots.font.size = Pt(10)
        r_dots.font.color.rgb = RGBColor(160, 160, 160)

        r_p = p.add_run(pg)
        r_p.font.name = "Times New Roman"
        r_p.font.size = Pt(11)
        r_p.font.italic = True
        r_p.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    # Abbreviations & List of Tables / Figures
    h_abbr = doc.add_heading("DANH MỤC CÁC TỪ VIẾT TẮT", level=1)
    h_abbr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h_abbr.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

    abbr_data = [
        ("SMB", "Server Message Block", "Giao thức truyền thông mạng chia sẻ tệp và máy in"),
        ("CIFS", "Common Internet File System", "Tên gọi khác của phương ngữ SMB phiên bản 1.0"),
        ("RCE", "Remote Code Execution", "Lỗ hổng thực thi mã tùy ý từ xa"),
        ("BSOD", "Blue Screen of Death", "Lỗi màn hình xanh dừng hệ điều hành Windows"),
        ("IPC", "Inter-Process Communication", "Cơ chế giao tiếp liên tiến trình qua đường ống ẩn IPC$"),
        ("RPC", "Remote Procedure Call", "Cơ chế gọi thủ tục từ xa trong kiến trúc mạng"),
        ("NBT", "NetBIOS over TCP/IP", "Phương thức truyền tải gói tin NetBIOS trên giao thức TCP/IP"),
        ("NSE", "Nmap Scripting Engine", "Cơ chế kịch bản tự động hóa mở rộng của công cụ Nmap"),
        ("MSF", "Metasploit Framework", "Bộ công cụ mã nguồn mở phục vụ kiểm thử xâm nhập"),
        ("FEA", "Full Extended Attribute", "Cấu trúc danh sách thuộc tính mở rộng đầy đủ trong hệ điều hành"),
        ("CVE", "Common Vulnerabilities and Exposures", "Hệ thống định danh chuẩn hóa các điểm yếu và lỗ hổng bảo mật"),
        ("VLAN", "Virtual Local Area Network", "Mạng cục bộ ảo phân tách miền quảng bá và định tuyến Layer 2"),
        ("UTM", "Unified Threat Management", "Thiết bị quản lý và phòng thủ mối đe dọa an ninh mạng thống nhất")
    ]

    t_abbr = doc.add_table(rows=len(abbr_data)+1, cols=3)
    t_abbr.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_abbr.autofit = False
    t_abbr.style = 'Table Grid'
    
    headers = ["Ký hiệu viết tắt", "Thuật ngữ tiếng Anh", "Ý nghĩa giải thích"]
    w_abbr = [Inches(1.2), Inches(2.3), Inches(2.7)]
    
    # Header row
    for idx, cell in enumerate(t_abbr.rows[0].cells):
        cell.width = w_abbr[idx]
        set_cell_background(cell, HEX_HEADER_BG)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(headers[idx])
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY

    # Data rows
    for row_i, (k, eng, vi) in enumerate(abbr_data):
        row = t_abbr.rows[row_i + 1]
        bg = HEX_ROW_ALT if row_i % 2 == 1 else "FFFFFF"
        for col_i, text in enumerate([k, eng, vi]):
            cell = row.cells[col_i]
            cell.width = w_abbr[col_i]
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)
            if col_i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.font.bold = True
                run.font.color.rgb = COLOR_NAVY
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # List of Tables
    h_lot = doc.add_heading("DANH MỤC BẢNG BIỂU", level=1)
    h_lot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h_lot.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

    tables_list = [
        ("Bảng 1.1", "So sánh đặc tính kỹ thuật cơ bản giữa SMBv1, SMBv2 và SMBv3", "Trang 4"),
        ("Bảng 1.2", "Các nhóm lỗ hổng SMB tiêu biểu trong bản tin an ninh MS17-010", "Trang 7"),
        ("Bảng 1.3", "Ánh xạ gói cập nhật KB4012212 theo các phiên bản Windows phổ biến", "Trang 10"),
        ("Bảng 1.4", "Khung đánh giá trạng thái an ninh giao thức SMB theo 4 cấp độ", "Trang 14"),
        ("Bảng 1.5", "Bảng đối chiếu trạng thái kỹ thuật và bề mặt rủi ro trước và sau phòng thủ", "Trang 16"),
        ("Bảng 2.1", "Ma trận phân bổ địa chỉ IP, VLAN và thông số kỹ thuật các nút mạng", "Trang 24"),
        ("Bảng 2.2", "Ma trận phân tích chuỗi nhân quả qua 4 trạng thái Snapshot đơn biến", "Trang 27"),
        ("Bảng 2.3", "Bảng đối chiếu các kỹ thuật kiểm thử với các tiền điều kiện và 4 cấp độ trạng thái", "Trang 35"),
        ("Bảng 2.4", "Ma trận bằng chứng thực nghiệm đối chiếu đa nguồn (Evidence Matrix)", "Trang 36")
    ]
    for b_id, b_name, b_pg in tables_list:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(b_id + ": ")
        r1.font.name = "Times New Roman"
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = COLOR_NAVY

        r2 = p.add_run(b_name)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11.5)

        dots = max(5, 70 - len(b_id + b_name))
        r_dots = p.add_run(" " + "·" * dots + " ")
        r_dots.font.color.rgb = RGBColor(160, 160, 160)

        r3 = p.add_run(b_pg)
        r3.font.name = "Times New Roman"
        r3.font.italic = True
        r3.font.size = Pt(11)
        r3.font.color.rgb = COLOR_MUTED

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # List of Figures
    h_lof = doc.add_heading("DANH MỤC HÌNH VÀ SƠ ĐỒ", level=1)
    h_lof.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h_lof.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_NAVY

    figures_list = [
        ("Sơ đồ 1.1", "Kiến trúc phân tầng Client – Server của dịch vụ SMB trên Windows", "Trang 2"),
        ("Sơ đồ 1.2", "Quy trình thiết lập phiên làm việc ba giai đoạn của giao thức SMB", "Trang 5"),
        ("Sơ đồ 1.3", "Cơ chế tràn số học danh sách FEA và ghi đè vùng nhớ Non-Paged Pool trong CVE-2017-0144", "Trang 9"),
        ("Sơ đồ 1.4", "Mô hình quan hệ phụ thuộc giữa 4 cấp độ đánh giá trạng thái an ninh SMB", "Trang 15"),
        ("Sơ đồ 2.1", "Sơ đồ topo mạng phân đoạn doanh nghiệp mô phỏng (Enterprise Network Topology)", "Trang 23"),
        ("Sơ đồ 2.2", "Sơ đồ tuần tự các pha kiểm thử theo mô hình chuỗi nhân quả", "Trang 35"),
        ("Sơ đồ 2.3", "Quy trình quản lý trạng thái máy ảo và cơ chế khôi phục tức thời qua Snapshot", "Trang 38")
    ]
    for s_id, s_name, s_pg in figures_list:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(s_id + ": ")
        r1.font.name = "Times New Roman"
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = COLOR_NAVY

        r2 = p.add_run(s_name)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11.5)

        dots = max(5, 70 - len(s_id + s_name))
        r_dots = p.add_run(" " + "·" * dots + " ")
        r_dots.font.color.rgb = RGBColor(160, 160, 160)

        r3 = p.add_run(s_pg)
        r3.font.name = "Times New Roman"
        r3.font.italic = True
        r3.font.size = Pt(11)
        r3.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

def add_formatted_text(p, text):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$|\$.*?\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            content = token[2:-2]
            run = p.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            content = token[1:-1]
            run = p.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            content = token[1:-1]
            run = p.add_run(content)
            run.font.name = 'Consolas'
            run.font.size = Pt(11)
            run.font.color.rgb = COLOR_NAVY
        elif token.startswith('$$') and token.endswith('$$') and len(token) >= 4:
            content = token[2:-2]
            run = p.add_run(content)
            run.font.name = 'Cambria Math'
            run.font.italic = True
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY
        elif token.startswith('$') and token.endswith('$') and len(token) >= 2:
            content = token[1:-1]
            run = p.add_run(content)
            run.font.name = 'Cambria Math'
            run.font.italic = True
            run.font.color.rgb = COLOR_NAVY
        else:
            run = p.add_run(token)
            run.font.name = 'Times New Roman'

def render_table(doc, table_lines):
    # Parse rows
    rows_data = []
    for line in table_lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line[1:-1].split('|')]
            # Check if it's separator row
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                continue
            rows_data.append(cells)
    
    if not rows_data or len(rows_data) < 2:
        return
    
    header = rows_data[0]
    data = rows_data[1:]
    num_cols = len(header)
    
    # Create docx table
    table = doc.add_table(rows=len(data) + 1, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = 'Table Grid'
    
    # Format header
    hdr_row = table.rows[0]
    for c_idx, cell_text in enumerate(header):
        cell = hdr_row.cells[c_idx]
        set_cell_background(cell, HEX_HEADER_BG)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_text(p, cell_text)
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = COLOR_NAVY

    # Format data rows
    for r_idx, row_vals in enumerate(data):
        row = table.rows[r_idx + 1]
        bg = HEX_ROW_ALT if r_idx % 2 == 1 else "FFFFFF"
        for c_idx in range(num_cols):
            cell = row.cells[c_idx]
            val = row_vals[c_idx] if c_idx < len(row_vals) else ""
            set_cell_background(cell, bg)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            # Text alignment heuristic
            if c_idx == 0 or len(val) <= 8 or re.match(r'^(True|False|\d+|0x[0-9A-Fa-f]+|S[0-3])$', val):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
            add_formatted_text(p, val)
            for r in p.runs:
                r.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def render_code_or_diagram(doc, block_lines):
    # Check if block is ASCII diagram
    text_content = "\n".join(block_lines)
    is_diagram = any(marker in text_content for marker in ['+---', '|   ', '+===', '====', '--->', '<---', '├───', '└───', '│'])
    
    # Create single-cell table for neat callout box
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = 'Table Grid'
    cell = table.cell(0, 0)
    
    bg_color = "F8FAFC" if is_diagram else HEX_CODE_BG
    set_cell_background(cell, bg_color)
    
    p = cell.paragraphs[0]
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    
    if is_diagram:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text_content)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text_content)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def parse_markdown_to_docx(doc, md_file_path):
    with open(md_file_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i]
        stripped = line.strip()
        
        # Skip empty lines
        if not stripped:
            i += 1
            continue
            
        # Horizontal rule
        if stripped == '---':
            i += 1
            continue
            
        # Code block / Diagram
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\r\n'))
                i += 1
            i += 1 # skip closing ```
            render_code_or_diagram(doc, code_lines)
            continue
            
        # Table
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            render_table(doc, table_lines)
            continue
            
        # Display Math ($$...$$)
        if stripped.startswith('$$') and stripped.endswith('$$') and len(stripped) > 4:
            math_text = stripped[2:-2].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(math_text)
            r.font.name = "Cambria Math"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        # Captions
        # Table Caption: *Bảng X.Y: ...*
        if stripped.startswith('*Bảng ') and stripped.endswith('*'):
            caption_text = stripped[1:-1]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(caption_text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        # Figure / Diagram Caption: *Sơ đồ X.Y: ...* or *Hình X.Y: ...*
        if (stripped.startswith('*Sơ đồ ') or stripped.startswith('*Hình ')) and stripped.endswith('*'):
            caption_text = stripped[1:-1]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(caption_text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11.5)
            r.font.italic = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        # Headings
        if stripped.startswith('# '):
            title = stripped[2:].strip()
            p = doc.add_heading(level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = "Times New Roman"
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        if stripped.startswith('## '):
            title = stripped[3:].strip()
            p = doc.add_heading(level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        if stripped.startswith('### '):
            title = stripped[4:].strip()
            p = doc.add_heading(level=3)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            i += 1
            continue

        if stripped.startswith('#### '):
            title = stripped[5:].strip()
            p = doc.add_heading(level=4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(title)
            r.font.name = "Times New Roman"
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = COLOR_DARK
            i += 1
            continue

        # Bullet List (- or *)
        if re.match(r'^[-*]\s+', stripped):
            item_text = re.sub(r'^[-*]\s+', '', stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            bullet_run = p.add_run("• ")
            bullet_run.font.name = "Times New Roman"
            bullet_run.font.bold = True
            bullet_run.font.color.rgb = COLOR_NAVY
            
            add_formatted_text(p, item_text)
            i += 1
            continue

        # Numbered List (1. 2. etc)
        m_num = re.match(r'^(\d+\.)\s+(.*)', stripped)
        if m_num:
            prefix = m_num.group(1)
            item_text = m_num.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            num_run = p.add_run(prefix + " ")
            num_run.font.name = "Times New Roman"
            num_run.font.bold = True
            num_run.font.color.rgb = COLOR_NAVY
            
            add_formatted_text(p, item_text)
            i += 1
            continue

        # Standard Paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.3
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Inches(0.4)
        add_formatted_text(p, stripped)
        i += 1

def build_complete_word_report():
    print("Initializing Document...")
    doc = setup_document()
    
    # 1. Front Cover Page
    print("Adding Front Cover Page...")
    add_cover_page(doc)
    
    # 2. Preliminary Section (TOC, Acronyms, LOT, LOF)
    print("Adding Preliminary Sections...")
    add_preliminary_section(doc)
    
    # 3. Main Body Section
    # New section with page numbers
    body_section = doc.add_section()
    body_section.page_width = Inches(8.27)
    body_section.page_height = Inches(11.69)
    body_section.top_margin = Inches(1.38)
    body_section.bottom_margin = Inches(1.18)
    body_section.left_margin = Inches(1.38)
    body_section.right_margin = Inches(0.79)
    add_header_footer_to_section(body_section, is_main_content=True)
    
    # Render Chapter 1
    ch1_path = "work/do-an/CHAPTER_1.md"
    print(f"Parsing and rendering {ch1_path}...")
    parse_markdown_to_docx(doc, ch1_path)
    
    # Page break between Chapter 1 and Chapter 2
    doc.add_page_break()
    
    # Render Chapter 2
    ch2_path = "work/do-an/CHAPTER_2.md"
    print(f"Parsing and rendering {ch2_path}...")
    parse_markdown_to_docx(doc, ch2_path)
    
    # Save output
    out_dir = "work/do-an"
    out_path = os.path.join(out_dir, "BAO_CAO_DO_AN_CHUONG_1_2.docx")
    print(f"Saving compiled report to {out_path}...")
    doc.save(out_path)
    print("SUCCESS: Document compiled successfully!")
    return out_path

if __name__ == "__main__":
    build_complete_word_report()
