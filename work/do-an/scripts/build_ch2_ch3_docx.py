import os
import re
import sys
import hashlib

sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Expected hashes and blobs
CH2_MD_PATH = 'work/do-an/CHAPTER_2.md'
CH3_MD_PATH = 'work/do-an/CHAPTER_3_DRAFT_R2.md'
CH2_FINAL_DOCX_PATH = 'work/do-an/output/CHAPTER_2_FINAL.docx'
OUTPUT_DOCX_PATH = 'work/do-an/output/CHAPTER_2_3_REVIEW.docx'

CH2_FINAL_SHA256 = '03adb6c44bc7c336b8c1be99f4e87c2ad93d0f11fe7d16603d5eb8dc102a7042'

COLOR_DARK = RGBColor(30, 41, 59)      # #1E293B Body text
COLOR_NAVY = RGBColor(31, 73, 125)     # #1F497D Navy accent
HEX_HEADER_BG = "E8EEF5"               # Soft ice blue / gray tint for table header

def compute_sha256(filepath):
    h = hashlib.sha256()
    with open(filepath, 'rb') as f:
        while chunk := f.read(8192):
            h.update(chunk)
    return h.hexdigest()

def verify_inputs():
    print("Verifying input files...")
    assert os.path.exists(CH2_MD_PATH), f"Missing {CH2_MD_PATH}"
    assert os.path.exists(CH3_MD_PATH), f"Missing {CH3_MD_PATH}"
    assert os.path.exists(CH2_FINAL_DOCX_PATH), f"Missing {CH2_FINAL_DOCX_PATH}"

    ch2_hash = compute_sha256(CH2_FINAL_DOCX_PATH)
    print(f"CHAPTER_2_FINAL.docx SHA-256: {ch2_hash}")
    assert ch2_hash == CH2_FINAL_SHA256, f"Hash mismatch: got {ch2_hash}, expected {CH2_FINAL_SHA256}"
    print("Verification PASSED!")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls("w"), fill_hex))
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>' % (nsdecls("w"), top, bottom, left, right))
    tcPr.append(tcMar)

def apply_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="B0B0B0"/>'
        r'  <w:insideV w:val="none"/>'
        r'  <w:left w:val="none"/>'
        r'  <w:right w:val="none"/>'
        r'</w:tblBorders>' % nsdecls("w")
    )
    tblPr.append(tblBorders)

def tokenize_md(text):
    """
    Parses markdown string with formatting into structured tokens:
    ('text', ...), ('bold', ...), ('italic', ...), ('code', ...), ('bold_code', ...), ('br', '\n')
    """
    text = re.sub(r'<br\s*/?>', '\n', text)
    tokens = []
    # Pattern to match bold_code, bold, italic, code, newline, or normal text
    pattern = re.compile(r'(\*\*`[^`]+`\*\*|\*\*.*?\*\*|\*[^*]+?\*|`[^`]+?`|\n|[^\*`\n]+)')
    pos = 0
    while pos < len(text):
        m = pattern.match(text, pos)
        if not m:
            tokens.append(('text', text[pos]))
            pos += 1
            continue
        chunk = m.group(1)
        pos = m.end()
        if chunk == '\n':
            tokens.append(('br', '\n'))
        elif chunk.startswith('**`') and chunk.endswith('`**'):
            tokens.append(('bold_code', chunk[3:-3]))
        elif chunk.startswith('**') and chunk.endswith('**'):
            tokens.append(('bold', chunk[2:-2]))
        elif chunk.startswith('*') and chunk.endswith('*') and len(chunk) > 2:
            tokens.append(('italic', chunk[1:-1]))
        elif chunk.startswith('`') and chunk.endswith('`'):
            tokens.append(('code', chunk[1:-1]))
        else:
            tokens.append(('text', chunk))
    return tokens

def add_formatted_runs_to_paragraph(p, text, base_font_name='Times New Roman', base_font_size=13.0, code_font_size=11.5, default_bold=False):
    tokens = tokenize_md(text)
    for t_type, t_val in tokens:
        if t_type == 'br':
            p.add_run('\n')
        elif t_type == 'bold':
            r = p.add_run(t_val)
            r.font.name = base_font_name
            r.font.size = Pt(base_font_size)
            r.font.bold = True
        elif t_type == 'italic':
            r = p.add_run(t_val)
            r.font.name = base_font_name
            r.font.size = Pt(base_font_size)
            r.font.italic = True
            if default_bold:
                r.font.bold = True
        elif t_type == 'code':
            r = p.add_run(t_val)
            r.font.name = 'Consolas'
            r.font.size = Pt(code_font_size)
            r.font.bold = True
        elif t_type == 'bold_code':
            r = p.add_run(t_val)
            r.font.name = 'Consolas'
            r.font.size = Pt(code_font_size)
            r.font.bold = True
        else:
            r = p.add_run(t_val)
            r.font.name = base_font_name
            r.font.size = Pt(base_font_size)
            if default_bold:
                r.font.bold = True

def add_body_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(28.35)  # 1.0 cm
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    add_formatted_runs_to_paragraph(p, text, base_font_name='Times New Roman', base_font_size=13.0, code_font_size=11.5)
    return p

def configure_heading_styles(doc):
    s1 = doc.styles['Heading 1']
    s1.font.name = 'Times New Roman'
    s1.font.size = Pt(16)
    s1.font.bold = True
    s1.font.color.rgb = RGBColor(0, 0, 0)
    s1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s1.paragraph_format.space_before = Pt(12)
    s1.paragraph_format.space_after = Pt(14)
    s1.paragraph_format.keep_with_next = True

    s2 = doc.styles['Heading 2']
    s2.font.name = 'Times New Roman'
    s2.font.size = Pt(14)
    s2.font.bold = True
    s2.font.color.rgb = RGBColor(0, 0, 0)
    s2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s2.paragraph_format.space_before = Pt(12)
    s2.paragraph_format.space_after = Pt(6)
    s2.paragraph_format.keep_with_next = True

    s3 = doc.styles['Heading 3']
    s3.font.name = 'Times New Roman'
    s3.font.size = Pt(13)
    s3.font.bold = True
    s3.font.italic = True
    s3.font.color.rgb = RGBColor(0, 0, 0)
    s3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s3.paragraph_format.space_before = Pt(8)
    s3.paragraph_format.space_after = Pt(4)
    s3.paragraph_format.keep_with_next = True

def align_ch2_heading_styles(doc):
    print("Aligning Chapter 2 heading paragraph styles to Heading 1/2/3...")
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith('CHƯƠNG 2.'):
            p.style = doc.styles['Heading 1']
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(14)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif re.match(r'^2\.\d+\.\s', t):
            p.style = doc.styles['Heading 2']
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)
        elif re.match(r'^2\.\d+\.\d+\.\s', t):
            p.style = doc.styles['Heading 3']
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.italic = True
                r.font.color.rgb = RGBColor(0, 0, 0)

def add_heading_1(doc, text):
    p = doc.add_paragraph(style=doc.styles['Heading 1'])
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph(style=doc.styles['Heading 2'])
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph(style=doc.styles['Heading 3'])
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_table_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    return p

def add_figure_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = False
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = True
    return p

def add_image_paragraph(doc, img_path, width_cm=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run()
    if width_cm:
        r.add_picture(img_path, width=Cm(width_cm))
    else:
        r.add_picture(img_path, width=Cm(14.8))
    return p

# Column width maps for Chapter 3 tables (in cm, summing to ~15.5 cm)
TABLE_WIDTHS = {
    'Bảng 3.1': [4.2, 4.8, 6.5],
    'Bảng 3.2': [4.5, 5.0, 6.0],
    'Bảng 3.3': [3.0, 3.8, 4.2, 4.5],
    'Bảng 3.4': [2.6, 4.2, 4.2, 4.5],
    'Bảng 3.5': [3.5, 3.0, 3.5, 5.5],
    'Bảng 3.6': [3.5, 3.0, 3.5, 5.5],
    'Bảng 3.7': [3.2, 4.1, 4.1, 4.1],
}

# Image width map for Chapter 3 figures (in cm)
IMAGE_WIDTHS = {
    'Hinh_3_1_Network_SMB.png': 14.5,
    'Hinh_3_2_Firewall.png': 14.5,
    'Hinh_3_3_SrvSys_Hotfix.png': 14.8,
    'Hinh_3_4_SMB_Version.png': 15.0,
    'Hinh_3_5_SMB_NSE.png': 14.8,
    'Hinh_3_6_MS17010_NSE.png': 15.0,
    'Hinh_3_7_After_Local.png': 14.8,
    'Hinh_3_8_SMB_Protocols_Retest.png': 15.0,
    'Hinh_3_9_pfSense_Rule_Order.png': 11.5,
    'Hinh_3_10_Nmap_Filtered.png': 12.5,
    'Hinh_3_11_pfSense_Block_Log.png': 15.2,
}

def create_word_table(doc, headers, data_rows, col_widths):
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    apply_table_borders(table)

    # Configure Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls("w")))
    trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls("w")))

    for c_idx, cell in enumerate(hdr_row.cells):
        cell.width = Cm(col_widths[c_idx])
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(headers[c_idx])
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.0)
        r.font.bold = True

    # Configure Data Rows
    for r_idx, row_data in enumerate(data_rows):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls("w")))

        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = Cm(col_widths[c_idx])
            set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            # Alignment: center for column 0 if short step/item name, or left
            if c_idx == 0 and len(cell_text) < 25 and not '\n' in cell_text and not '<br>' in cell_text:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_runs_to_paragraph(p, cell_text, base_font_name='Times New Roman', base_font_size=9.5, code_font_size=8.0)

    # Add a clean space after table
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)
    p_spacer.paragraph_format.line_spacing = 1.0
    r_empty = p_spacer.add_run()
    r_empty.font.size = Pt(2)
    return table

def parse_and_append_chapter_3(doc):
    print("Parsing Chapter 3 markdown and appending to document...")
    with open(CH3_MD_PATH, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    total_lines = len(lines)
    current_table_title = None

    while i < total_lines:
        line = lines[i].strip()

        if not line:
            i += 1
            continue

        # Heading 1
        if line.startswith('# '):
            text = line[2:].strip()
            add_heading_1(doc, text)
            i += 1
            continue

        # Heading 2
        if line.startswith('## '):
            text = line[3:].strip()
            add_heading_2(doc, text)
            i += 1
            continue

        # Heading 3
        if line.startswith('### '):
            text = line[4:].strip()
            add_heading_3(doc, text)
            i += 1
            continue

        # Table title: **Bảng 3.x. ...**
        if line.startswith('**Bảng 3.') and line.endswith('**'):
            raw_title = line[2:-2].strip()
            add_table_title(doc, raw_title)
            # Find table key e.g. 'Bảng 3.1'
            m_key = re.match(r'(Bảng 3\.\d+)', raw_title)
            current_table_title = m_key.group(1) if m_key else 'Bảng 3.1'
            i += 1
            continue

        # Table rows: starts with |
        if line.startswith('|') and ('---' not in line):
            # Parse entire markdown table
            table_lines = []
            while i < total_lines and lines[i].strip().startswith('|'):
                l_str = lines[i].strip()
                if not re.match(r'^\|[-:| ]+\|$', l_str):
                    table_lines.append(l_str)
                i += 1

            headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            data_rows = []
            for r_line in table_lines[1:]:
                cells = [c.strip() for c in r_line.split('|')[1:-1]]
                data_rows.append(cells)

            col_widths = TABLE_WIDTHS.get(current_table_title, [15.5 / len(headers)] * len(headers))
            create_word_table(doc, headers, data_rows, col_widths)
            continue

        # Image
        if line.startswith('!['):
            # Form 1: ![Hình 3.x. Caption text](path)
            # Form 2: ![Hình 3.x](path) followed by *Hình 3.x. Caption text*
            m_img = re.match(r'!\[(.*?)\]\((.*?)\)', line)
            if m_img:
                alt_text = m_img.group(1).strip()
                rel_path = m_img.group(2).strip()
                full_img_path = os.path.join('work/do-an', rel_path)

                img_filename = os.path.basename(rel_path)
                width_cm = IMAGE_WIDTHS.get(img_filename, 14.8)

                # Check if caption is in alt_text or on next line
                caption_text = None
                if alt_text.startswith('Hình 3.') and '.' in alt_text[8:]:
                    caption_text = alt_text

                # Check next line for caption
                if not caption_text and i + 1 < total_lines:
                    next_line = lines[i+1].strip()
                    if next_line.startswith('*Hình 3.') and next_line.endswith('*'):
                        caption_text = next_line[1:-1].strip()
                        i += 1 # Consume caption line

                add_image_paragraph(doc, full_img_path, width_cm=width_cm)
                if caption_text:
                    add_figure_caption(doc, caption_text)
                else:
                    add_figure_caption(doc, alt_text)
            i += 1
            continue

        # Standalone caption line (if not consumed above)
        if line.startswith('*Hình 3.') and line.endswith('*'):
            add_figure_caption(doc, line[1:-1].strip())
            i += 1
            continue

        # Regular prose paragraph
        add_body_paragraph(doc, line)
        i += 1

def build_combined_document():
    verify_inputs()

    print("Loading base Chapter 2 document...")
    doc = Document(CH2_FINAL_DOCX_PATH)

    # Configure true Word heading styles and align Chapter 2 headings
    configure_heading_styles(doc)
    align_ch2_heading_styles(doc)

    print("Adding page break for Chapter 3...")
    doc.add_page_break()

    parse_and_append_chapter_3(doc)

    print(f"Saving combined document to {OUTPUT_DOCX_PATH}...")
    os.makedirs(os.path.dirname(OUTPUT_DOCX_PATH), exist_ok=True)
    doc.save(OUTPUT_DOCX_PATH)
    print("Combined document saved successfully!")

    # Audit structure
    audit_document(doc)

def audit_document(doc):
    print("\n--- STRUCTURAL AUDIT ---")
    h1_list = []
    h2_list = []
    h3_list = []
    tables_list = []
    figures_list = []

    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t.startswith('CHƯƠNG'):
            h1_list.append(t)
        elif re.match(r'^[23]\.\d+\.\s', t):
            h2_list.append(t)
        elif re.match(r'^[23]\.\d+\.\d+\.\s', t):
            h3_list.append(t)
        elif re.match(r'^Bảng\s+[23]\.\d+\.', t):
            tables_list.append(t)
        elif re.match(r'^Hình\s+[23]\.\d+\.', t):
            figures_list.append(t)

    print(f"Heading 1 ({len(h1_list)}): {h1_list}")
    print(f"Heading 2 ({len(h2_list)}):")
    for h in h2_list:
        print(f"  {h}")
    print(f"Heading 3 ({len(h3_list)}):")
    for h in h3_list:
        print(f"  {h}")
    print(f"Tables ({len(tables_list)}):")
    for tbl in tables_list:
        print(f"  {tbl}")
    print(f"Figures ({len(figures_list)}):")
    for fig in figures_list:
        print(f"  {fig}")

    assert len(h1_list) == 2, f"Expected 2 H1, got {len(h1_list)}"
    assert len(h2_list) == 14, f"Expected 14 H2, got {len(h2_list)}"
    assert len(h3_list) == 30, f"Expected 30 H3, got {len(h3_list)}"
    assert len(tables_list) == 11, f"Expected 11 tables, got {len(tables_list)}"
    assert len(figures_list) == 13, f"Expected 13 figures, got {len(figures_list)}"

    print("\n--- HEADING STYLE & OUTLINE LEVEL AUDIT ---")
    h1_style_list = [p for p in doc.paragraphs if p.style and (p.style.name == 'Heading 1' or p.style.style_id == 'Heading1')]
    h2_style_list = [p for p in doc.paragraphs if p.style and (p.style.name == 'Heading 2' or p.style.style_id == 'Heading2')]
    h3_style_list = [p for p in doc.paragraphs if p.style and (p.style.name == 'Heading 3' or p.style.style_id == 'Heading3')]

    print(f"Heading 1 by style ({len(h1_style_list)}): {[p.text[:35] for p in h1_style_list]}")
    print(f"Heading 2 by style ({len(h2_style_list)}): {[p.text[:35] for p in h2_style_list]}")
    print(f"Heading 3 by style ({len(h3_style_list)}): {[p.text[:35] for p in h3_style_list]}")

    assert len(h1_style_list) == 2, f"Expected 2 Heading 1 by style, got {len(h1_style_list)}"
    assert len(h2_style_list) == 14, f"Expected 14 Heading 2 by style, got {len(h2_style_list)}"
    assert len(h3_style_list) == 30, f"Expected 30 Heading 3 by style, got {len(h3_style_list)}"

    print("\n--- SECTION MARGIN AUDIT ---")
    sec = doc.sections[0]
    top_cm = sec.top_margin.cm
    bottom_cm = sec.bottom_margin.cm
    left_cm = sec.left_margin.cm
    right_cm = sec.right_margin.cm
    w_cm = sec.page_width.cm
    h_cm = sec.page_height.cm
    print(f"Section 0: top={top_cm:.2f} cm, bottom={bottom_cm:.2f} cm, left={left_cm:.2f} cm, right={right_cm:.2f} cm, width={w_cm:.2f} cm, height={h_cm:.2f} cm")
    assert abs(top_cm - 3.5) < 0.05, f"Top margin mismatch: {top_cm} != 3.5 cm"
    assert abs(bottom_cm - 3.0) < 0.05, f"Bottom margin mismatch: {bottom_cm} != 3.0 cm"
    assert abs(left_cm - 3.5) < 0.05, f"Left margin mismatch: {left_cm} != 3.5 cm"
    assert abs(right_cm - 2.0) < 0.05, f"Right margin mismatch: {right_cm} != 2.0 cm"
    assert abs(w_cm - 21.0) < 0.05, f"Page width mismatch: {w_cm} != 21.0 cm"
    assert abs(h_cm - 29.7) < 0.05, f"Page height mismatch: {h_cm} != 29.7 cm"

    print("\n--- TECHNICAL TRUTH LOCK AUDIT ---")
    full_text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for t in doc.tables for r in t.rows for cell in r.cells])

    required_phrases = [
        "445 OPEN != vulnerable",
        "SMBv1 enabled != MS17-010 confirmed",
        "SMBv1 disabled != PATCHED",
        "FILTERED != PATCHED",
        "UNKNOWN != SAFE"
    ]
    for rp in required_phrases:
        assert rp in full_text, f"Missing required truth phrase: '{rp}'"
        print(f"  [PASS] Found required phrase: '{rp}'")

    forbidden_phrases = [
        "Bảng 3.8",
        "Hình 3.12",
        "CHƯƠNG 1.",
        "CHƯƠNG 4.",
        "meterpreter",
        "reverse_tcp",
        "exploit/windows/smb/ms17_010",
        "định tuyến qua pfSense",
        "bảo đảm chỉ mở cho trạm kiểm thử chỉ định"
    ]
    for fp in forbidden_phrases:
        assert fp not in full_text, f"Found forbidden phrase: '{fp}'"
        print(f"  [PASS] Verified absence of forbidden phrase: '{fp}'")

    print("\nALL STRUCTURAL, STYLE, MARGIN & TRUTH AUDIT ASSERTIONS PASSED!")

if __name__ == '__main__':
    build_combined_document()
