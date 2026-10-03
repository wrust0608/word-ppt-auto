import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colors
COLOR_NAVY = RGBColor(0, 32, 96)         # #002060 UTC Primary Navy
COLOR_BLUE = RGBColor(31, 78, 121)       # #1F4E79 Secondary
COLOR_DARK = RGBColor(38, 38, 38)        # #262626 Body text
COLOR_MUTED = RGBColor(89, 89, 89)       # #595959 Secondary
COLOR_RED = RGBColor(192, 0, 0)          # #C00000 Alert
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(176, 196, 222)   # #B0C4DE Border

HEX_HEADER_BG = "EBF1F5"                 # Soft elegant UTC ice blue
HEX_ROW_ALT = "F9FBFC"
HEX_BORDER = "B0C4DE"
HEX_CALLOUT_BG = "F4F6F9"

def format_table_cell(cell, fill_hex=None, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith(('shd', 'tcMar')):
            tcPr.remove(child)
    if fill_hex:
        shd = parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls("w"), fill_hex))
        tcPr.append(shd)
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>' % (nsdecls("w"), top, left, bottom, right))
    tcPr.append(tcMar)

def set_table_borders(table, color_hex="D3D3D3"):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="single" w:sz="6" w:space="0" w:color="%s"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="002060"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>' % (nsdecls('w'), color_hex, color_hex)
    )
    inserted = False
    for child in list(tblPr):
        if child.tag.endswith(('tblLayout', 'tblCellMar', 'tblLook')):
            child.addprevious(tblBorders)
            inserted = True
            break
    if not inserted:
        tblPr.append(tblBorders)

def add_footer_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_header_footer(section, is_main=False):
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    
    # Header
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.text = ""
    if is_main:
        hrun = hp.add_run("TRƯỜNG ĐH GIAO THÔNG VẬN TẢI (UTC) • QUẢN LÝ DỰ ÁN ĐẦU TƯ XÂY DỰNG")
        hrun.font.name = "Times New Roman"
        hrun.font.size = Pt(8.5)
        hrun.font.italic = True
        hrun.font.color.rgb = COLOR_MUTED

    # Footer
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.text = ""
    if is_main:
        frun = fp.add_run("Trang ")
        frun.font.name = "Times New Roman"
        frun.font.size = Pt(9.5)
        frun.font.color.rgb = COLOR_MUTED
        add_footer_page_number(frun)

def setup_document():
    doc = Document()
    
    # Configure A4 & UTC Margins: Top 2.0cm, Bottom 2.0cm, Left 3.0cm, Right 2.0cm
    section = doc.sections[0]
    section.page_width = Inches(8.27)     # 210mm
    section.page_height = Inches(11.69)   # 297mm
    section.top_margin = Inches(0.79)     # 20mm
    section.bottom_margin = Inches(0.79)  # 20mm
    section.left_margin = Inches(1.18)    # 30mm
    section.right_margin = Inches(0.79)   # 20mm
    
    # Set default style to Times New Roman
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = COLOR_DARK
    style_normal.paragraph_format.line_spacing = 1.4
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.space_before = Pt(0)
    
    return doc

def add_cover_page(doc):
    # Front cover section
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    r1 = p.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(13)
    r1.font.bold = True
    
    r2 = p.add_run("TRƯỜNG ĐẠI HỌC GIAO THÔNG VẬN TẢI HÀ NỘI\n")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_NAVY
    
    r3 = p.add_run("KHOA KỸ THUẬT XÂY DỰNG / KINH TẾ XÂY DỰNG\n")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True
    
    # Horizontal separator
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(40)
    r_div = p_div.add_run("―" * 25)
    r_div.font.color.rgb = COLOR_NAVY
    
    # Title Box
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run("TIỂU LUẬN CHUYÊN ĐỀ HỌC THUẬT")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_MUTED
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(14)
    r_t = p_title.add_run("QUẢN LÝ DỰ ÁN ĐẦU TƯ XÂY DỰNG CÔNG TRÌNH")
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(19)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_NAVY
    
    p_theme = doc.add_paragraph()
    p_theme.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_theme.paragraph_format.space_after = Pt(50)
    r_theme = p_theme.add_run("CHUYÊN ĐỀ: BỘ BÀI LÀM VÀ BÁO CÁO PHẢN BIỆN ĐA TẦNG 5 CÂU HỎI TRỌNG TÂM\n(Phân tích nhiệm vụ, Vòng đời dự án, Phối hợp chủ thể,\nXử lý mất đồng bộ Cung ứng và Giải quyết tranh chấp chậm tiến độ 3 Nhà thầu)")
    r_theme.font.name = "Times New Roman"
    r_theme.font.size = Pt(12.5)
    r_theme.font.italic = True
    r_theme.font.color.rgb = COLOR_DARK
    
    # Meta Info Table
    t = doc.add_table(rows=4, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    
    labels = [
        ("Học phần / Môn học:", "Quản lý dự án đầu tư xây dựng công trình"),
        ("Mã học phần / Lớp:", "Chuyên ngành Xây dựng Công trình Giao thông & Dân dụng"),
        ("Mục tiêu đánh giá:", "Đạt điểm tối đa (10/10) - Đạt chuẩn phản biện đa tầng UTC"),
        ("Năm học / Học kỳ:", "Năm học 2025 - 2026")
    ]
    
    for idx, (label, val) in enumerate(labels):
        row = t.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.0)
        
        p_l = cell_lbl.paragraphs[0]
        p_l.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_l.paragraph_format.space_after = Pt(3)
        rl = p_l.add_run(label)
        rl.font.name = "Times New Roman"
        rl.font.size = Pt(12)
        rl.font.bold = True
        rl.font.color.rgb = COLOR_NAVY
        
        p_v = cell_val.paragraphs[0]
        p_v.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_v.paragraph_format.space_after = Pt(3)
        rv = p_v.add_run(val)
        rv.font.name = "Times New Roman"
        rv.font.size = Pt(12)
        rv.font.color.rgb = COLOR_DARK
        
        format_table_cell(cell_lbl, fill_hex=None, top=40, bottom=40, left=60, right=60)
        format_table_cell(cell_val, fill_hex=None, top=40, bottom=40, left=60, right=60)

    # Footer Cover
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot.paragraph_format.space_before = Pt(80)
    p_foot.paragraph_format.space_after = Pt(0)
    rf = p_foot.add_run("HÀ NỘI, NĂM 2026")
    rf.font.name = "Times New Roman"
    rf.font.size = Pt(12)
    rf.font.bold = True
    rf.font.color.rgb = COLOR_NAVY
    
    doc.add_page_break()

def parse_inline_markdown(paragraph, text, default_font_size=Pt(13), default_color=COLOR_DARK, is_bold_default=False, is_italic_default=False):
    # Parses bold (**text**), italic (*text*), code (`text`)
    pattern = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = paragraph.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.size = default_font_size
            run.font.bold = True
            run.font.italic = is_italic_default
            run.font.color.rgb = default_color
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            content = part[1:-1]
            run = paragraph.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.size = default_font_size
            run.font.bold = is_bold_default
            run.font.italic = True
            run.font.color.rgb = default_color
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            run = paragraph.add_run(content)
            run.font.name = 'Consolas'
            run.font.size = Pt(default_font_size.pt - 1)
            run.font.bold = is_bold_default
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = default_font_size
            run.font.bold = is_bold_default
            run.font.italic = is_italic_default
            run.font.color.rgb = default_color

def render_markdown_table(doc, table_lines):
    rows_data = []
    for line in table_lines:
        line = line.strip()
        if not line.startswith('|') or not line.endswith('|'):
            continue
        # Check if separator row
        inner = line[1:-1].strip()
        if all(c in '-: |' for c in inner):
            continue
        cells = [c.strip() for c in line[1:-1].split('|')]
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    t = doc.add_table(rows=len(rows_data), cols=num_cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    
    # Calculate column widths
    avail_width = 6.3  # Inches
    col_width = avail_width / num_cols
    
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = t.cell(r_idx, c_idx)
            cell.width = Inches(col_width)
            val = row[c_idx] if c_idx < len(row) else ""
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            if is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                format_table_cell(cell, fill_hex=HEX_HEADER_BG, top=60, bottom=60, left=100, right=100)
                parse_inline_markdown(p, val, default_font_size=Pt(11), default_color=COLOR_NAVY, is_bold_default=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                bg = HEX_ROW_ALT if (r_idx % 2 == 1) else None
                format_table_cell(cell, fill_hex=bg, top=60, bottom=60, left=100, right=100)
                parse_inline_markdown(p, val, default_font_size=Pt(11), default_color=COLOR_DARK)
                
    set_table_borders(t, HEX_BORDER)
    # Add space after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_callout_box(doc, text):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = Inches(6.3)
    
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith(('tcBorders', 'shd', 'tcMar')):
            tcPr.remove(child)
            
    tcBorders = parse_xml(
        r'<w:tcBorders %s>'
        r'<w:top w:val="none"/>'
        r'<w:left w:val="single" w:sz="24" w:space="0" w:color="002060"/>'
        r'<w:bottom w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'</w:tcBorders>' % nsdecls('w')
    )
    shd = parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls("w"), HEX_CALLOUT_BG))
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="100" w:type="dxa"/><w:left w:w="150" w:type="dxa"/><w:bottom w:w="100" w:type="dxa"/><w:right w:w="150" w:type="dxa"/></w:tcMar>' % nsdecls("w"))
    
    tcPr.append(tcBorders)
    tcPr.append(shd)
    tcPr.append(tcMar)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.3
    parse_inline_markdown(p, text, default_font_size=Pt(11.5), default_color=COLOR_DARK, is_italic_default=True)
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)

def convert_markdown_file_to_docx(doc, md_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_table = False
    table_buffer = []
    in_code_block = False
    code_buffer = []
    
    for line in lines:
        raw_line = line.rstrip('\r\n')
        stripped = raw_line.strip()
        
        # Handle code block
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                # render code block
                code_text = "\n".join(code_buffer)
                p_c = doc.add_paragraph()
                p_c.paragraph_format.space_before = Pt(4)
                p_c.paragraph_format.space_after = Pt(6)
                p_c.paragraph_format.left_indent = Inches(0.2)
                p_c.paragraph_format.line_spacing = 1.1
                r_c = p_c.add_run(code_text)
                r_c.font.name = "Consolas"
                r_c.font.size = Pt(10)
                r_c.font.color.rgb = RGBColor(40, 40, 40)
                code_buffer = []
            else:
                in_code_block = True
                code_buffer = []
            continue
            
        if in_code_block:
            code_buffer.append(raw_line)
            continue
            
        # Handle table
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_buffer.append(stripped)
            continue
        else:
            if in_table:
                in_table = False
                render_markdown_table(doc, table_buffer)
                table_buffer = []
                
        # Empty line
        if not stripped:
            continue
            
        # Callout quote
        if stripped.startswith('>'):
            callout_text = stripped.lstrip('>').strip()
            add_callout_box(doc, callout_text)
            continue
            
        # Headings
        if stripped.startswith('# '):
            heading_text = stripped[2:].strip()
            doc.add_page_break()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(10)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            continue
            
        if stripped.startswith('## '):
            heading_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            continue
            
        if stripped.startswith('### '):
            heading_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = COLOR_BLUE
            continue
            
        if stripped.startswith('#### '):
            heading_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = COLOR_DARK
            continue
            
        # Horizontal rule
        if stripped.startswith('---'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("―" * 35)
            r.font.color.rgb = COLOR_BORDER
            continue
            
        # Bullet list
        if stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('+ '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.left_indent = Inches(0.25)
            
            # Bullet symbol
            r_b = p.add_run("•  ")
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(11)
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_NAVY
            
            parse_inline_markdown(p, bullet_text, default_font_size=Pt(12.5), default_color=COLOR_DARK)
            continue
            
        # Numbered list or indented bullet
        if re.match(r'^\d+\.\s+', stripped):
            m = re.match(r'^(\d+\.)\s+(.*)$', stripped)
            num_prefix = m.group(1)
            rest_text = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.left_indent = Inches(0.25)
            
            r_num = p.add_run(f"{num_prefix} ")
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(12.5)
            r_num.font.bold = True
            r_num.font.color.rgb = COLOR_NAVY
            
            parse_inline_markdown(p, rest_text, default_font_size=Pt(12.5), default_color=COLOR_DARK)
            continue
            
        # Sub-bullet (e.g. "  + ")
        if raw_line.startswith('    +') or raw_line.startswith('  +') or raw_line.startswith('    *'):
            sub_text = stripped.lstrip('+* ').strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.left_indent = Inches(0.45)
            
            r_b = p.add_run("―  ")
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = COLOR_MUTED
            
            parse_inline_markdown(p, sub_text, default_font_size=Pt(12), default_color=COLOR_DARK)
            continue
            
        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.4
        p.paragraph_format.first_line_indent = Inches(0.3)
        parse_inline_markdown(p, stripped, default_font_size=Pt(13), default_color=COLOR_DARK)
        
    if in_table:
        render_markdown_table(doc, table_buffer)

def main():
    base_dir = r"e:\word_ppt-auto\work\quan-ly-du-an-xaydung"
    out_docx = os.path.join(base_dir, "BAO_CAO_BAI_LAM_QLDA_XD.docx")
    answers_md = os.path.join(base_dir, "DE_AN_TRA_LOI_CAU_HOI.md")
    review_md = os.path.join(base_dir, "REVIEW_REPORT_PHAN_BIEN.md")
    
    print("Khởi tạo tài liệu Word theo chuẩn Trường ĐH Giao thông Vận tải Hà Nội...")
    doc = setup_document()
    
    print("Tạo trang bìa chuẩn mẫu UTC...")
    add_cover_page(doc)
    
    # Configure main section headers and footers
    section_main = doc.sections[0]
    add_header_footer(section_main, is_main=True)
    
    print(f"Xử lý nội dung bài làm từ {answers_md}...")
    convert_markdown_file_to_docx(doc, answers_md)
    
    print(f"Xử lý báo cáo phản biện đa tầng từ {review_md}...")
    convert_markdown_file_to_docx(doc, review_md)
    
    print(f"Lưu file kết quả tại: {out_docx}")
    doc.save(out_docx)
    print("Hoàn tất tạo file DOCX thành công!")

if __name__ == "__main__":
    main()
