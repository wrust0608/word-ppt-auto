import os
import re
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

COLOR_NAVY = RGBColor(0, 32, 96)         # #002060 UTC Primary Navy
COLOR_BLUE = RGBColor(31, 78, 121)       # #1F4E79 Secondary
COLOR_DARK = RGBColor(38, 38, 38)        # #262626 Body text
COLOR_MUTED = RGBColor(89, 89, 89)       # #595959 Secondary
COLOR_BORDER = RGBColor(176, 196, 222)   # #B0C4DE Border

HEX_HEADER_BG = "EBF1F5"
HEX_ROW_ALT = "F9FBFC"
HEX_BORDER = "B0C4DE"

def format_table_cell(cell, fill_hex=None, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith(('shd', 'tcMar')):
            tcPr.remove(child)
    if fill_hex:
        shd = parse_xml(r'<w:shd %s w:val="clear" w:color="auto" w:fill="%s"/>' % (nsdecls("w"), fill_hex))
        tcPr.append(shd)
    tcMar = parse_xml(r'<w:tcMar %s><w:top w:w="%d" w:type="dxa"/><w:left w:w="%d" w:type="dxa"/><w:bottom w:w="%d" w:type="dxa"/><w:right w:w="%d" w:type="dxa"/></w:tcMar>' % (nsdecls("w"), top, left, bottom, right))
    tcPr.append(tcMar)

def set_table_borders(table, color_hex="D3D3D3"):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'<w:top w:val="single" w:sz="6" w:space="0" w:color="%s"/>'
        r'<w:left w:val="none"/>'
        r'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="002060"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="%s"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>' % (nsdecls('w'), color_hex, color_hex)
    )
    inserted = False
    for child in list(tblPr):
        if child.tag.endswith(('tblLayout', 'tblCellMar', 'tblLook')):
            child.addprevious(tblBorders)
            inserted = True
            break
    if not inserted:
        tblPr.append(tblBorders)

def add_footer_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)     # 210mm
    section.page_height = Inches(11.69)   # 297mm
    section.top_margin = Inches(0.79)     # 20mm
    section.bottom_margin = Inches(0.79)  # 20mm
    section.left_margin = Inches(1.18)    # 30mm
    section.right_margin = Inches(0.79)   # 20mm
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.4)
    
    # Footer simple page number
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.text = ""
    frun = fp.add_run("Trang ")
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(9.5)
    frun.font.color.rgb = COLOR_MUTED
    add_footer_page_number(frun)
    
    # Default Normal style
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = COLOR_DARK
    style_normal.paragraph_format.line_spacing = 1.35
    style_normal.paragraph_format.space_after = Pt(4)
    style_normal.paragraph_format.space_before = Pt(0)
    
    return doc

def parse_inline_markdown(paragraph, text, default_font_size=Pt(13), default_color=COLOR_DARK, is_bold_default=False, is_italic_default=False):
    pattern = re.compile(r'(\*\*[^*]+?\*\*|\*[^*]+?\*|`[^`]+?`)')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = paragraph.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.size = default_font_size
            run.font.bold = True
            run.font.italic = is_italic_default
            run.font.color.rgb = default_color
        elif part.startswith('*') and part.endswith('*') and not part.startswith('**'):
            content = part[1:-1]
            run = paragraph.add_run(content)
            run.font.name = 'Times New Roman'
            run.font.size = default_font_size
            run.font.bold = is_bold_default
            run.font.italic = True
            run.font.color.rgb = default_color
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            run = paragraph.add_run(content)
            run.font.name = 'Consolas'
            run.font.size = Pt(default_font_size.pt - 1)
            run.font.bold = is_bold_default
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            run = paragraph.add_run(part)
            run.font.name = 'Times New Roman'
            run.font.size = default_font_size
            run.font.bold = is_bold_default
            run.font.italic = is_italic_default
            run.font.color.rgb = default_color

def render_markdown_table(doc, table_lines):
    rows_data = []
    for line in table_lines:
        line = line.strip()
        if not line.startswith('|') or not line.endswith('|'):
            continue
        inner = line[1:-1].strip()
        if all(c in '-: |' for c in inner):
            continue
        cells = [c.strip() for c in line[1:-1].split('|')]
        rows_data.append(cells)
        
    if not rows_data:
        return
        
    num_cols = max(len(r) for r in rows_data)
    t = doc.add_table(rows=len(rows_data), cols=num_cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    
    avail_width = 6.3
    if num_cols >= 6:
        col_widths = [1.6] + [(avail_width - 1.6) / (num_cols - 1)] * (num_cols - 1)
    elif num_cols > 2:
        col_widths = [1.8] + [(avail_width - 1.8) / (num_cols - 1)] * (num_cols - 1)
    else:
        col_widths = [avail_width / num_cols] * num_cols
    
    for r_idx, row in enumerate(rows_data):
        is_header = (r_idx == 0)
        for c_idx in range(num_cols):
            cell = t.cell(r_idx, c_idx)
            cell.width = Inches(col_widths[c_idx])
            val = row[c_idx] if c_idx < len(row) else ""
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
            if is_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                format_table_cell(cell, fill_hex=HEX_HEADER_BG, top=60, bottom=60, left=50, right=50)
                parse_inline_markdown(p, val, default_font_size=Pt(10), default_color=COLOR_NAVY, is_bold_default=True)
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
                bg = HEX_ROW_ALT if (r_idx % 2 == 1) else None
                format_table_cell(cell, fill_hex=bg, top=60, bottom=60, left=50, right=50)
                parse_inline_markdown(p, val, default_font_size=Pt(9.5), default_color=COLOR_DARK)
                
    set_table_borders(t, HEX_BORDER)
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def build_simple_qa_docx(md_path=None, out_docx=None):
    if md_path is None:
        md_path = sys.argv[1] if len(sys.argv) > 1 else r"e:\word_ppt-auto\work\quan-ly-du-an-xaydung\TRA_LOI_CAU_HOI_QLDA_BAN_NOP.md"
    if out_docx is None:
        out_docx = sys.argv[2] if len(sys.argv) > 2 else r"e:\word_ppt-auto\work\quan-ly-du-an-xaydung\TRA_LOI_CAU_HOI_QLDA_BAN_NOP.docx"
    
    doc = setup_document()
    
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_table = False
    table_buffer = []
    in_code_block = False
    code_buffer = []
    
    for line in lines:
        raw_line = line.rstrip('\r\n')
        stripped = raw_line.strip()
        
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                code_text = "\n".join(code_buffer)
                p_c = doc.add_paragraph()
                p_c.paragraph_format.space_before = Pt(4)
                p_c.paragraph_format.space_after = Pt(6)
                p_c.paragraph_format.left_indent = Inches(0.2)
                p_c.paragraph_format.line_spacing = 1.1
                r_c = p_c.add_run(code_text)
                r_c.font.name = "Consolas"
                r_c.font.size = Pt(10)
                r_c.font.color.rgb = RGBColor(40, 40, 40)
                code_buffer = []
            else:
                in_code_block = True
                code_buffer = []
            continue
            
        if in_code_block:
            code_buffer.append(raw_line)
            continue
            
        if stripped.startswith('|') and stripped.endswith('|'):
            in_table = True
            table_buffer.append(stripped)
            continue
        else:
            if in_table:
                in_table = False
                render_markdown_table(doc, table_buffer)
                table_buffer = []
                
        if not stripped:
            continue
            
        # Top Header (H1)
        if stripped.startswith('# '):
            heading_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            continue
            
        # Question Headings (H2)
        if stripped.startswith('## '):
            heading_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13.5)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY
            continue
            
        # Sub Headings (H3)
        if stripped.startswith('### '):
            heading_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = COLOR_BLUE
            continue
            
        # H4
        if stripped.startswith('#### '):
            heading_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(heading_text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12.5)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = COLOR_DARK
            continue
            
        if stripped.startswith('---'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run("―" * 35)
            r.font.color.rgb = COLOR_BORDER
            continue
            
        if stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('+ '):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.left_indent = Inches(0.25)
            
            r_b = p.add_run("•  ")
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(11)
            r_b.font.bold = True
            r_b.font.color.rgb = COLOR_NAVY
            
            parse_inline_markdown(p, bullet_text, default_font_size=Pt(12.5), default_color=COLOR_DARK)
            continue
            
        if re.match(r'^\d+\.\s+', stripped):
            m = re.match(r'^(\d+\.)\s+(.*)$', stripped)
            num_prefix = m.group(1)
            rest_text = m.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.left_indent = Inches(0.25)
            
            r_num = p.add_run(f"{num_prefix} ")
            r_num.font.name = "Times New Roman"
            r_num.font.size = Pt(12.5)
            r_num.font.bold = True
            r_num.font.color.rgb = COLOR_NAVY
            
            parse_inline_markdown(p, rest_text, default_font_size=Pt(12.5), default_color=COLOR_DARK)
            continue
            
        if raw_line.startswith('    +') or raw_line.startswith('  +') or raw_line.startswith('    *'):
            sub_text = stripped.lstrip('+* ').strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.left_indent = Inches(0.45)
            
            r_b = p.add_run("―  ")
            r_b.font.name = "Times New Roman"
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = COLOR_MUTED
            
            parse_inline_markdown(p, sub_text, default_font_size=Pt(12), default_color=COLOR_DARK)
            continue
            
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.first_line_indent = Inches(0.25)
        parse_inline_markdown(p, stripped, default_font_size=Pt(13), default_color=COLOR_DARK)
        
    if in_table:
        render_markdown_table(doc, table_buffer)
        
    try:
        doc.save(out_docx)
        print("Đã tạo file Word thành công:", out_docx)
    except PermissionError:
        alt_docx = out_docx.replace(".docx", "_moinhat.docx")
        doc.save(alt_docx)
        print(f"File {out_docx} đang mở trong Word, đã lưu thành công bản mới nhất vào: {alt_docx}")

if __name__ == "__main__":
    build_simple_qa_docx()
