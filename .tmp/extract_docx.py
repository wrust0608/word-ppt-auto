from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def clean(value: str) -> str:
    return " ".join(value.replace("\u00a0", " ").split())


def extract(path: Path) -> str:
    doc = Document(path)
    lines = [f"# {path.name}", "", "## Paragraphs", ""]
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = clean(paragraph.text)
        if text:
            lines.append(f"P{index:04d} [{paragraph.style.name}]: {text}")

    lines.extend(["", "## Tables", ""])
    for table_index, table in enumerate(doc.tables, start=1):
        lines.append(f"### Table {table_index}")
        for row_index, row in enumerate(table.rows, start=1):
            values = [clean(cell.text) for cell in row.cells]
            lines.append(f"R{row_index:03d}: " + " | ".join(values))
        lines.append("")
    return "\n".join(lines)


if __name__ == "__main__":
    source = Path(sys.argv[1])
    destination = Path(sys.argv[2])
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(extract(source), encoding="utf-8")
