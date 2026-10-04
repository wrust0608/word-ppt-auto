"""Build the reviewed Chapter 1 DOCX using semantic Word styles and HUIT rules."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def set_font(style, size=13, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.append(fonts)
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Times New Roman')
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)
    color = rpr.find(qn('w:color'))
    if color is not None:
        for key in ['themeColor', 'themeTint', 'themeShade']:
            color.attrib.pop(qn('w:' + key), None)


def add_formatted_runs(p, text):
    """Parse basic markdown bold, italic, and code spans into Word runs."""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
        else:
            run = p.add_run(token)
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0, 0, 0)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3.5), Cm(3.0)
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(2.0)
    sec.footer_distance = Cm(1.5)

    # Base paragraph styling
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            set_font(style)
            fmt = style.paragraph_format
            fmt.line_spacing = 1.5
            fmt.space_before, fmt.space_after = Pt(0), Pt(6)
            fmt.widow_control = True
            ppr = style.element.find(qn('w:pPr'))
            if ppr is not None:
                for border in list(ppr.findall(qn('w:pBdr'))):
                    ppr.remove(border)

    # Normal style (HUIT 2024: Times New Roman 13pt, 1.5 line spacing, 1.25cm first line indent, justified)
    normal = doc.styles['Normal']
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    # Headings
    # Heading 1: Chapter title (14pt, bold, centered, keep_with_next)
    h1 = doc.styles['Heading 1']
    set_font(h1, 14, True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.line_spacing = 1.5
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True

    # Heading 2: Section title (13pt, bold, left, keep_with_next)
    h2 = doc.styles['Heading 2']
    set_font(h2, 13, True)
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.first_line_indent = Cm(0)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.line_spacing = 1.5
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    # Heading 3: Subsection title (13pt, bold, left, keep_with_next)
    h3 = doc.styles['Heading 3']
    set_font(h3, 13, True)
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.first_line_indent = Cm(0)
    h3.paragraph_format.space_before = Pt(6)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.line_spacing = 1.5
    h3.paragraph_format.keep_with_next = True
    h3.paragraph_format.keep_together = True

    # Footer page numbering
    footer = doc.styles['Footer']
    set_font(footer, 10, False)
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    fp = sec.footer.paragraphs[0]
    fp.style = footer
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = '1'
    run.append(text)
    field.append(run)
    fp._p.append(field)

    number = OxmlElement('w:pgNumType')
    number.set(qn('w:fmt'), 'decimal')
    number.set(qn('w:start'), '1')
    sec._sectPr.insert_element_before(number, 'w:cols', 'w:docGrid')

    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.insert_element_before(
        update, 'w:compat', 'w:docVars', 'w:rsids', 'w:mathPr', 'w:themeFontLang', 'w:clrSchemeMapping'
    )

    doc.core_properties.title = 'Chương 1 — Cơ sở lý thuyết và công cụ kiểm thử SMB'
    doc.core_properties.subject = 'Đồ án chuyên ngành An toàn thông tin'

    source = ROOT / 'work/do-an/CHAPTER_1.md'
    lines = source.read_text(encoding='utf-8').splitlines()

    for line in lines:
        stripped = line.strip()
        if not stripped or stripped == '---':
            continue

        if line.startswith('# '):
            p = doc.add_paragraph(style='Heading 1')
            add_formatted_runs(p, line[2:].strip())
        elif line.startswith('## '):
            p = doc.add_paragraph(style='Heading 2')
            add_formatted_runs(p, line[3:].strip())
        elif line.startswith('### '):
            p = doc.add_paragraph(style='Heading 3')
            add_formatted_runs(p, line[4:].strip())
        elif line.startswith('  - ') or line.startswith('    - '):
            # Sub-bullet
            content = line.strip()[2:].strip()
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.75)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            bullet_run = p.add_run('– ')
            bullet_run.font.name = 'Times New Roman'
            bullet_run.font.color.rgb = RGBColor(0, 0, 0)
            add_formatted_runs(p, content)
        elif line.startswith('- '):
            # Top-level bullet
            content = line[2:].strip()
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            bullet_run = p.add_run('• ')
            bullet_run.font.name = 'Times New Roman'
            bullet_run.font.color.rgb = RGBColor(0, 0, 0)
            add_formatted_runs(p, content)
        else:
            # Regular paragraph
            p = doc.add_paragraph(style='Normal')
            # Check if this paragraph is a lead-in to a list
            if stripped.endswith(':'):
                p.paragraph_format.keep_with_next = True
            add_formatted_runs(p, stripped)

    out = ROOT / 'work/do-an/outputs/CHUONG_1_REVIEW.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f'Saved Chapter 1 review docx to: {out}')


if __name__ == '__main__':
    build()
