"""Build the reviewed Introduction using semantic Word styles and HUIT rules."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def set_font(style, size=13, bold=False):
    style.font.name = 'Times New Roman'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Times New Roman')
    for key in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        fonts.attrib.pop(qn('w:' + key), None)
    color = rpr.find(qn('w:color'))
    for key in ['themeColor', 'themeTint', 'themeShade']:
        color.attrib.pop(qn('w:' + key), None)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3.5), Cm(3)
    sec.left_margin, sec.right_margin = Cm(3.5), Cm(2)
    sec.footer_distance = Cm(1.5)
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            set_font(style)
            fmt = style.paragraph_format
            fmt.line_spacing = 1.5
            fmt.space_before, fmt.space_after = Pt(0), Pt(6)
            fmt.widow_control = True
            ppr = style.element.find(qn('w:pPr'))
            if ppr is not None:
                for border in list(ppr.findall(qn('w:pBdr'))):
                    ppr.remove(border)
    normal = doc.styles['Normal']
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    for name, size, before, after in [('Title',16,0,9),('Heading 1',14,9,6),('Heading 2',13,6,3)]:
        style = doc.styles[name]
        set_font(style, size, True)
        fmt = style.paragraph_format
        fmt.first_line_indent = Cm(0)
        fmt.space_before, fmt.space_after = Pt(before), Pt(after)
        fmt.keep_with_next = True
        fmt.keep_together = True
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == 'Title' else WD_ALIGN_PARAGRAPH.LEFT
    bib = doc.styles.add_style('Bibliography HUIT', WD_STYLE_TYPE.PARAGRAPH)
    bib.base_style = normal
    set_font(bib)
    bf = bib.paragraph_format
    bf.left_indent, bf.first_line_indent = Cm(0.8), Cm(-0.8)
    bf.line_spacing = 1.5
    bf.space_before, bf.space_after = Pt(0), Pt(3)
    bf.keep_together, bf.widow_control = True, True
    bf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = doc.styles['Footer']
    set_font(footer)
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    fp = sec.footer.paragraphs[0]
    fp.style = footer
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' PAGE ')
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = '1'
    run.append(text)
    field.append(run)
    fp._p.append(field)
    number = OxmlElement('w:pgNumType')
    number.set(qn('w:fmt'), 'decimal')
    number.set(qn('w:start'), '1')
    sec._sectPr.insert_element_before(number, 'w:cols', 'w:docGrid')
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    doc.settings.element.insert_element_before(update, 'w:compat', 'w:docVars', 'w:rsids', 'w:mathPr', 'w:themeFontLang', 'w:clrSchemeMapping')
    doc.core_properties.title = 'Mở đầu — Xây dựng mô hình kiểm thử lỗ hổng SMB trên Windows bằng Kali Linux'
    doc.core_properties.subject = 'Đồ án chuyên ngành An toàn thông tin'
    source = ROOT / 'work/do-an/INTRODUCTION.md'
    for line in source.read_text(encoding='utf-8').splitlines():
        if not line.strip():
            continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], style='Title')
        elif line.startswith('## '):
            p = doc.add_paragraph(line[3:], style='Heading 1')
            if line[3:] == 'TÀI LIỆU THAM KHẢO':
                p.paragraph_format.page_break_before = True
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], style='Heading 2')
        else:
            style = 'Bibliography HUIT' if re.match(r'^\[\d+\]\s', line) else 'Normal'
            doc.add_paragraph(line, style=style)
    out = ROOT / 'work/do-an/outputs/MO_DAU_DO_AN.docx'
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == '__main__':
    build()
